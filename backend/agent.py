import os
import sys
import sqlite3
import json
import time
import re
import argparse
import urllib.request
import urllib.error
import urllib.parse
import email as email_module
from datetime import datetime, timedelta, timezone

# Ensure robust UTF-8 printing on Windows console
try:
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    if hasattr(sys.stderr, 'reconfigure'):
        sys.stderr.reconfigure(encoding='utf-8', errors='replace')
except Exception:
    pass


# Optional dependencies
try:
    from imapclient import IMAPClient
    IMAP_AVAILABLE = True
except ImportError:
    IMAP_AVAILABLE = False

try:
    import openpyxl
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False

try:
    from cryptography.hazmat.primitives.ciphers import Cipher, algorithms, modes
    from cryptography.hazmat.backends import default_backend
    import hashlib
    CRYPTOGRAPHY_AVAILABLE = True
except ImportError:
    CRYPTOGRAPHY_AVAILABLE = False

DB_FILE = os.path.join(os.path.dirname(__file__), '..', 'users.db')
EXCEL_FILE = os.path.join(os.path.dirname(__file__), '..', 'data', 'exports', 'logs.xlsx')

# ── Mock Emails for Simulation ───────────────────────────────────────────────

MOCK_EMAILS = [
    {
        "sender": "sarah.jenkins@techcorp.com",
        "subject": "Quote Request: Custom Web Application Development",
        "body": "Hi team,\n\nWe are looking to develop a customer portal for our retail branch. We require features like auth, payment gateways, and a beautiful analytics dashboard. Our budget is around $25,000, and we need this ready in 3 months. Can you send us a quote?\n\nBest regards,\nSarah Jenkins\nDirector of Operations, TechCorp",
        "priority": "Medium",
        "category": "important",
        "intent": "Request for project quote and timeline discussion",
        "summary": "Sarah Jenkins from TechCorp requested a quote for a custom retail customer portal. Budget is $25k, deadline is 3 months.",
        "requires_approval": 1,
        "action_items": [
            "Prepare a development estimation sheet.",
            "Draft response email asking for API specification details.",
            "Schedule an introductory call with Sarah."
        ],
        "draft_reply": "Subject: Re: Quote Request: Custom Web Application Development\n\nDear Sarah,\n\nThank you for reaching out to us. We would be delighted to assist TechCorp in developing your retail customer portal. Your budget of $25,000 and timeline of 3 months align with our capabilities.\n\nTo provide an accurate quote, could you share any detailed API specifications or design guidelines you have? I'd also love to schedule a brief 15-minute introductory call next Tuesday to align on key details. Please let me know your availability.\n\nBest regards,\nSales & Account Team\nAI Email Automation"
    },
    {
        "sender": "emma.watson@talentpartners.io",
        "subject": "Interview Invitation: Lead AI Engineer Position",
        "body": "Hi,\n\nI was very impressed by your profile and experience in building full-stack AI automation workflows. Our hiring committee would love to invite you for a 45-minute technical conversation this Friday at 3:00 PM EST.\n\nPlease confirm if this time slot works for you or suggest an alternative that fits your calendar.\n\nWarm regards,\nEmma Watson\nSenior Technical Recruiter",
        "priority": "High",
        "category": "important",
        "intent": "Interview scheduling & career opportunity discussion",
        "summary": "Emma Watson from Talent Partners invited you to a 45-minute technical interview for Lead AI Engineer this Friday at 3:00 PM EST.",
        "requires_approval": 1,
        "action_items": [
            "Confirm Friday 3:00 PM EST interview availability.",
            "Review portfolio and system architecture materials.",
            "Send updated resume copy."
        ],
        "draft_reply": "Subject: Re: Interview Invitation: Lead AI Engineer Position\n\nHi Emma,\n\nThank you for reaching out! I would be thrilled to speak with the hiring committee. Friday at 3:00 PM EST works perfectly for me.\n\nLooking forward to our conversation.\n\nBest regards"
    },
    {
        "sender": "david.miller@devcraft.org",
        "subject": "Quick question regarding API Integration milestone",
        "body": "Hey there,\n\nHope you're having a productive week! Just checking in to see if we can sync on the OAuth 2.0 and IMAP webhook integration endpoints tomorrow morning. Let me know what time works best for a quick 10-minute huddle.\n\nCheers,\nDavid Miller\nLead Software Architect",
        "priority": "Medium",
        "category": "important",
        "intent": "Project collaboration & technical sync check-in",
        "summary": "David Miller asked for a 10-minute sync tomorrow morning regarding OAuth 2.0 and IMAP webhook integration endpoints.",
        "requires_approval": 1,
        "action_items": [
            "Check calendar for tomorrow morning opening.",
            "Prepare quick status update on IMAP webhook endpoints."
        ],
        "draft_reply": "Subject: Re: Quick question regarding API Integration milestone\n\nHi David,\n\nTomorrow morning works great. How does 10:30 AM sound? I will send over a calendar invite with the meeting link.\n\nBest,\nTeam"
    },
    {
        "sender": "monitoring@cloudalert.net",
        "subject": "URGENT: Production Server Database Load > 95%",
        "body": "ALERT: database-prod-01 is experiencing critically high CPU and RAM usage (98.4%). Connection latency has increased to 4500ms. Response failures are occurring on main API endpoints.\n\nTimestamp: 2026-07-19T11:45:00Z\nHost: aws-us-east-1",
        "priority": "High",
        "category": "important",
        "intent": "Critical infrastructure alert requiring immediate action",
        "summary": "Production database database-prod-01 is overloaded (98.4% usage), leading to API response failures and 4500ms latencies.",
        "requires_approval": 0,
        "action_items": [
            "Check database active connection pool.",
            "Identify and kill any slow-running queries.",
            "Notify the DevOps on-call engineer."
        ],
        "draft_reply": "Subject: Re: URGENT: Production Server Database Load > 95%\n\nHello Team,\n\nThis is an automated acknowledgment: We have received the database load alert for 'database-prod-01'. Priority has been elevated to CRITICAL.\n\nOur on-call engineers have been paged to identify slow queries and optimize connection pools immediately. Updates will follow on the status channel within 10 minutes.\n\nBest regards,\nSupport Operations"
    },
    {
        "sender": "newsletter@marketinginsights.org",
        "subject": "Weekly Tech Trends Newsletter - July 2026",
        "body": "Welcome to your weekly digest! This week, we explore the rise of Node.js 22 SQLite integrations, local LLM configurations with Ollama, and CSS grid layout tips.",
        "priority": "Low",
        "category": "newsletter",
        "intent": "Marketing newsletter — no action required",
        "summary": "Weekly tech newsletter covering Node.js 22 SQLite, local LLM setups with Ollama, and CSS grid tips.",
        "requires_approval": 0,
        "action_items": [],
        "draft_reply": ""
    },
    {
        "sender": "promo@deals-store.com",
        "subject": "🔥 Flash Sale: 70% Off All Electronics!",
        "body": "Don't miss our biggest flash sale of the year! Get up to 70% off on all electronics. Limited time only. Shop now at deals-store.com.",
        "priority": "Low",
        "category": "promotion",
        "intent": "Promotional advertisement — no action required",
        "summary": "Flash sale promotional email from deals-store.com offering 70% off electronics.",
        "requires_approval": 0,
        "action_items": [],
        "draft_reply": ""
    }
]

# ── Environment Loader ───────────────────────────────────────────────────────

def load_env():
    """Load .env from project root."""
    env_path = os.path.join(os.path.dirname(__file__), '..', '.env')
    if not os.path.exists(env_path):
        return
    with open(env_path, 'r', encoding='utf-8') as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith('#'):
                continue
            eq = line.find('=')
            if eq > 0:
                key = line[:eq].strip()
                val = line[eq+1:].strip()
                os.environ[key] = val

load_env()

# ── Decryption Helper ────────────────────────────────────────────────────────

def decrypt_password(encrypted_text):
    if not encrypted_text:
        return ""
    if ":" not in encrypted_text:
        return encrypted_text  # Fallback to plaintext if not encrypted
    if not CRYPTOGRAPHY_AVAILABLE:
        print("[Warn] Cryptography module not available. Cannot decrypt password.")
        return encrypted_text
    try:
        secret_key = os.environ.get("ENCRYPTION_KEY", "")
        if not secret_key:
            return encrypted_text
        
        iv_hex, cipher_hex = encrypted_text.split(":", 1)
        iv = bytes.fromhex(iv_hex)
        ciphertext = bytes.fromhex(cipher_hex)
        
        if len(secret_key) == 64:
            try:
                key_bytes = bytes.fromhex(secret_key)
            except ValueError:
                key_bytes = hashlib.sha256(secret_key.encode('utf-8')).digest()
        else:
            key_bytes = hashlib.sha256(secret_key.encode('utf-8')).digest()
            
        cipher = Cipher(algorithms.AES(key_bytes), modes.CBC(iv), backend=default_backend())
        decryptor = cipher.decryptor()
        decrypted_padded = decryptor.update(ciphertext) + decryptor.finalize()
        
        padding_len = decrypted_padded[-1]
        if padding_len < 1 or padding_len > 16:
            return decrypted_padded.decode('utf-8')
        decrypted = decrypted_padded[:-padding_len]
        return decrypted.decode('utf-8')
    except Exception as e:
        print(f"[Decrypt Error] Failed to decrypt: {e}")
        return encrypted_text

# ── Gemini API with Rate-Limit Retry ─────────────────────────────────────────

def generate_with_gemini(api_key, sender, subject, body):
    """Calls Gemini API with exponential backoff on 429 errors."""
    model = os.environ.get("GEMINI_MODEL", "gemini-flash-lite-latest")
    url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent?key={api_key}"

    prompt = f"""You are an intelligent email assistant. Analyze this email and return a JSON response with these fields:
- "priority": "High", "Medium", or "Low"
- "category": classify as one of: "important", "spam", "newsletter", "promotion" (Company, internship, interview, meeting, invitation, and business client emails MUST be classified as "important")
- "intent": A brief phrase describing the sender's intent (e.g., "Interview Invitation", "Internship Opportunity", "Meeting Request", "Company Offer", "Quote Request", "Newsletter")
- "summary": A 2-3 sentence clear summary
- "action_items": An array of action item strings (empty array if none)
- "draft_reply": A professional, polite, and context-aware reply draft (or empty string if no reply needed)
- "requires_approval": true or false. Set to true if the email is an important communication requiring user review/approval before sending, specifically:
  1. Company, corporate, or official business emails
  2. Internship opportunities, job offers, or career inquiries
  3. Interview invitations or scheduling requests
  4. Meetings, appointments, calls, calendar invites, or date/time coordination
  5. Event invitations or client contract proposals
  Set to false only for generic promotional ads, routine automated newsletters, or low-priority alerts that do not need personal user approval.

Email:
From: {sender}
Subject: {subject}
Body:
{body}
"""

    data = {
        "contents": [{"parts": [{"text": prompt}]}],
        "generationConfig": {"responseMimeType": "application/json"}
    }

    req_body = json.dumps(data).encode('utf-8')

    max_attempts = 5
    base_sleep = 5

    for attempt in range(1, max_attempts + 1):
        req = urllib.request.Request(url, data=req_body, headers={'Content-Type': 'application/json'})
        try:
            if attempt > 1:
                print(f"[Gemini API] Retry attempt {attempt}/{max_attempts}...")
            else:
                print(f"[Gemini API] Analyzing email from {sender}...")

            with urllib.request.urlopen(req, timeout=8) as response:
                res_data = json.loads(response.read().decode('utf-8'))
                text = res_data['candidates'][0]['content']['parts'][0]['text'].strip()
                if text.startswith('```json'):
                    text = text[7:]
                if text.endswith('```'):
                    text = text[:-3]
                return json.loads(text.strip())

        except urllib.error.HTTPError as e:
            error_body = e.read().decode('utf-8')
            print(f"[Gemini API Error] HTTP {e.code}: {error_body[:200]}")

            if e.code == 429:
                sleep_time = base_sleep * (2 ** (attempt - 1))
                match = re.search(r"Please retry in ([0-9.]+)s", error_body)
                if match:
                    suggested = float(match.group(1))
                    sleep_time = max(sleep_time, suggested + 1.0)
                print(f"[Rate Limited] Sleeping {sleep_time:.1f}s before retrying...")
                time.sleep(sleep_time)
            else:
                return None
    return None

def heuristic_fallback(sender, subject, body):
    """Robust local heuristic classifier used if Gemini API is unreachable or rate-limited."""
    sender_l = (sender or "").lower()
    subj_l = (subject or "").lower()
    body_l = (body or "").lower()
    combined = f"{sender_l} {subj_l} {body_l}"

    # Priority & Category Heuristics
    if any(k in combined for k in ["interview", "job", "hiring", "offer", "candidate", "internship", "career"]):
        category = "important"
        priority = "High"
        intent = "Career & Interview Opportunity"
        requires_approval = True
    elif any(k in combined for k in ["meeting", "schedule", "call", "urgent", "invoice", "payment", "quote", "client", "proposal"]):
        category = "important"
        priority = "High"
        intent = "Business Coordination"
        requires_approval = True
    elif any(k in combined for k in ["deal", "discount", "sale", "promo", "coupon", "offer"]):
        category = "promotion"
        priority = "Low"
        intent = "Promotional Offer"
        requires_approval = False
    elif any(k in combined for k in ["newsletter", "weekly", "digest", "update", "roundup", "edition"]):
        category = "newsletter"
        priority = "Low"
        intent = "Newsletter Update"
        requires_approval = False
    elif any(k in combined for k in ["spam", "lottery", "winner", "prize", "crypto"]):
        category = "spam"
        priority = "Low"
        intent = "Unsolicited Message"
        requires_approval = False
    else:
        category = "important"
        priority = "Medium"
        intent = "General Communication"
        requires_approval = True

    # Clean summary
    clean_body = re.sub(r'\s+', ' ', body).strip()
    if clean_body and clean_body != "Empty body.":
        summary = clean_body[:200] + ("..." if len(clean_body) > 200 else "")
    else:
        summary = f"Direct email from {sender} with subject '{subject}'."

    # Action items
    action_items = []
    if "interview" in combined or "job" in combined:
        action_items.append("Review interview schedule and confirm availability")
        action_items.append("Prepare required portfolio / resume documents")
    elif "meeting" in combined:
        action_items.append("Check calendar and confirm meeting time slot")
    elif "invoice" in combined or "payment" in combined:
        action_items.append("Verify payment terms and billing invoice")
    elif category == "important":
        action_items.append("Follow up with sender regarding message")

    # Smart Draft reply
    if category == "important":
        name_part = sender.split("@")[0].replace(".", " ").replace("_", " ").title()
        draft_reply = (
            f"Hi {name_part},\n\n"
            f"Thank you for reaching out regarding \"{subject}\". I have received your message and will review the details and get back to you shortly.\n\n"
            f"Best regards,\n[Your Name]"
        )
    else:
        draft_reply = ""

    return {
        "priority": priority,
        "category": category,
        "intent": intent,
        "summary": summary,
        "action_items": action_items,
        "draft_reply": draft_reply,
        "requires_approval": requires_approval
    }


# ── Excel Logging ────────────────────────────────────────────────────────────

def initialize_excel():
    if not EXCEL_AVAILABLE:
        print("[Warn] openpyxl not installed. Excel logging skipped. Run: pip install openpyxl")
        return
    os.makedirs(os.path.dirname(EXCEL_FILE), exist_ok=True)
    if not os.path.exists(EXCEL_FILE):
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Email Logs"
            ws.append(["Timestamp", "Email ID", "Sender", "Subject", "Priority", "Category", "Intent", "Status"])
            wb.save(EXCEL_FILE)
            print(f"[OK] Created Excel logs: {EXCEL_FILE}")
        except Exception as e:
            print(f"[Error] Excel init failed: {e}")

def log_to_excel(email_id, sender, subject, priority, category, intent, status):
    if not EXCEL_AVAILABLE:
        return
    try:
        try:
            wb = openpyxl.load_workbook(EXCEL_FILE)
            ws = wb.active
        except Exception:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Email Logs"
            ws.append(["Timestamp", "Email ID", "Sender", "Subject", "Priority", "Category", "Intent", "Status"])
        ws.append([datetime.now().strftime("%Y-%m-%d %H:%M:%S"), email_id, sender, subject, priority, category, intent, status])
        wb.save(EXCEL_FILE)
        print(f"[Excel] Logged email {email_id}")
    except Exception as e:
        print(f"[Excel Error] {e}")


def get_existing_email_id(sender, subject, body, user_id=1):
    try:
        conn = sqlite3.connect(DB_FILE, timeout=10.0)
        c = conn.cursor()
        body_snippet = body[:100] if body else ""
        c.execute("SELECT id FROM emails WHERE user_id = ? AND sender = ? AND subject = ? AND body LIKE ?", (user_id, sender, subject, body_snippet + '%'))
        row = c.fetchone()
        conn.close()
        return row[0] if row else None
    except Exception as e:
        print(f"[DB Error checking existence] {e}")
        return None

def get_existing_email_id_fast(sender, subject, user_id=1):
    try:
        conn = sqlite3.connect(DB_FILE, timeout=5.0)
        c = conn.cursor()
        c.execute("SELECT id FROM emails WHERE user_id = ? AND sender = ? AND subject = ? LIMIT 1", (user_id, sender, subject))
        row = c.fetchone()
        conn.close()
        return row[0] if row else None
    except Exception:
        return None

# ── Database Insert ──────────────────────────────────────────────────────────

def insert_email_to_db(email, user_id=1, imap_account_id=None):
    try:
        conn = sqlite3.connect(DB_FILE, timeout=10.0)
        c = conn.cursor()
        
        # Ensure column exists
        try:
            c.execute("ALTER TABLE emails ADD COLUMN requires_approval INTEGER DEFAULT 0")
            conn.commit()
        except sqlite3.OperationalError:
            pass
        try:
            c.execute("ALTER TABLE emails ADD COLUMN imap_account_id INTEGER")
            conn.commit()
        except sqlite3.OperationalError:
            pass
            
        received_at = datetime.now(timezone.utc).isoformat().replace('+00:00', 'Z')
        
        # Check auto-reply policy rules
        db_settings = load_db_settings()
        master_auto_reply = db_settings.get("rule_master_autoreply", "true").lower() != "false"
        delay_mins_str = db_settings.get("rule_delay_mins", "20")
        try:
            delay_mins = int(delay_mins_str)
        except ValueError:
            delay_mins = 20

        category = (email.get("category") or "important").lower()
        
        # Check category-specific permission (default: Important=ON, Promo/Newsletter/Spam=OFF)
        category_allowed = True
        if category == "important":
            category_allowed = db_settings.get("rule_reply_important", "true").lower() != "false"
        elif category == "promotion":
            category_allowed = db_settings.get("rule_reply_promotion", "false").lower() == "true"
        elif category == "newsletter":
            category_allowed = db_settings.get("rule_reply_newsletter", "false").lower() == "true"
        elif category == "spam":
            category_allowed = db_settings.get("rule_reply_spam", "false").lower() == "true"
            
        requires_approval = email.get("requires_approval", 0)

        # Set auto-send schedule timer if allowed
        if not master_auto_reply or not category_allowed or requires_approval or delay_mins <= 0 or not email.get("draft_reply"):
            send_at = None
        else:
            send_at = (datetime.now(timezone.utc) + timedelta(minutes=delay_mins)).isoformat().replace('+00:00', 'Z')
            
        action_items_str = json.dumps(email.get("action_items", []))

        c.execute('''
            INSERT INTO emails (user_id, sender, subject, body, received_at, priority, status, category, intent, summary, action_items, draft_reply, send_at, requires_approval, imap_account_id)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (
            user_id,
            email["sender"],
            email["subject"],
            email.get("body", ""),
            received_at,
            email.get("priority", "Medium"),
            "Pending",
            email.get("category", "important"),
            email.get("intent", ""),
            email.get("summary", ""),
            action_items_str,
            email.get("draft_reply", ""),
            send_at,
            requires_approval,
            imap_account_id
        ))

        eid = c.lastrowid

        # Insert action items as tasks
        action_items = email.get("action_items", [])
        for item in action_items:
            c.execute('''
                INSERT INTO tasks (user_id, email_id, title, description, status, priority, created_at)
                VALUES (?, ?, ?, ?, ?, ?, ?)
            ''', (user_id, eid, item, f"Action item extracted from email: '{email['subject']}'", 'To Do', email.get("priority", "Medium"), received_at))

        conn.commit()
        conn.close()
        print(f"[DB] Inserted email {eid} from {email['sender']} [Category: {email.get('category', 'unknown')}] and created {len(action_items)} tasks")
        return eid
    except Exception as e:
        print(f"[DB Error] {e}")
        return None

# ── Simulation Mode ──────────────────────────────────────────────────────────

def run_simulation(user_id=1):
    print(f"\n--- Running AI Email Agent Simulation for User ID {user_id} ---")
    initialize_excel()
    inserted_count = 0
    for mock in MOCK_EMAILS:
        existing = get_existing_email_id(mock["sender"], mock["subject"], mock.get("body", ""), user_id)
        if existing:
            print(f"[Skip] Mock email '{mock['subject']}' already exists for user {user_id} (ID: {existing})")
            continue
        print(f"\n[Ingesting] Email from: {mock['sender']} [{mock['category']}] for user {user_id}")
        time.sleep(0.3)
        eid = insert_email_to_db(mock, user_id=user_id)
        if eid:
            inserted_count += 1
            log_to_excel(eid, mock["sender"], mock["subject"], mock["priority"], mock.get("category", ""), mock.get("intent", ""), "Pending")
    print(f"\n[OK] Simulation complete. {inserted_count} new emails ingested for user {user_id}.")


# ── Google OAuth Token Refresh ───────────────────────────────────────────────

GOOGLE_CLIENT_ID = os.environ.get('GOOGLE_CLIENT_ID', '')
GOOGLE_CLIENT_SECRET = os.environ.get('GOOGLE_CLIENT_SECRET', '')

def refresh_google_access_token(refresh_token):
    """Exchange a refresh token for a fresh access token."""
    try:
        load_env()
        client_id = os.environ.get('GOOGLE_CLIENT_ID') or GOOGLE_CLIENT_ID
        client_secret = os.environ.get('GOOGLE_CLIENT_SECRET') or GOOGLE_CLIENT_SECRET
        data = urllib.parse.urlencode({
            'client_id': client_id,
            'client_secret': client_secret,
            'refresh_token': refresh_token,
            'grant_type': 'refresh_token'
        }).encode('utf-8')
        
        req = urllib.request.Request('https://oauth2.googleapis.com/token', data=data, method='POST')
        req.add_header('Content-Type', 'application/x-www-form-urlencoded')
        
        with urllib.request.urlopen(req, timeout=15) as response:
            result = json.loads(response.read().decode('utf-8'))
            access_token = result.get('access_token')
            if access_token:
                print(f"[OAuth] Successfully refreshed access token")
                return access_token
            else:
                print(f"[OAuth Error] No access_token in response: {result}")
                return None
    except Exception as e:
        print(f"[OAuth Error] Failed to refresh token: {e}")
        return None

# ── Real IMAP Agent ──────────────────────────────────────────────────────────

def run_real_agent(imap_server, imap_user, imap_pass, llm_api_key, user_id=1, imap_account_id=None, auth_type='password', oauth_refresh_token=None):
    if not IMAP_AVAILABLE:
        print("[Error] imapclient not installed. Run: pip install imapclient")
        return
    if not llm_api_key or llm_api_key == "your_gemini_api_key_here":
        print("[Error] Valid GEMINI_API_KEY required. Check settings page or .env.")
        return

    print(f"\n--- Starting AI Email Agent for {imap_user} ({imap_server}) [auth: {auth_type}] ---")
    initialize_excel()

    try:
        print(f"[IMAP] Connecting to {imap_server}...")
        with IMAPClient(imap_server, ssl=True) as client:
            if auth_type == 'oauth' and oauth_refresh_token:
                # Get fresh access token using refresh token
                access_token = refresh_google_access_token(oauth_refresh_token)
                if not access_token:
                    print(f"[Error] Failed to refresh OAuth access token for {imap_user}")
                    return
                # Authenticate with XOAUTH2
                oauth2_string = f"user={imap_user}\x01auth=Bearer {access_token}\x01\x01"
                client.oauth2_login(imap_user, access_token)
                print(f"[IMAP] OAuth2 login successful for {imap_user}!")
            else:
                client.login(imap_user, imap_pass)
                print(f"[IMAP] Login successful for {imap_user}!")
            client.select_folder('INBOX')

            # Strategy: Check UNSEEN first, then also grab today's emails (SINCE today)
            # This ensures emails read on mobile (marked Seen) are still captured.
            from datetime import date as _date
            today_str = _date.today().strftime('%d-%b-%Y')

            unseen_messages = client.search('UNSEEN')
            today_messages = client.search(['SINCE', today_str])

            # Combine both lists (unique message IDs)
            combined = list(set(list(unseen_messages) + list(today_messages)))
            # Sort so newest IDs (highest number) are last, take last 15 max
            combined.sort()
            if len(combined) > 15:
                combined = combined[-15:]

            if combined:
                messages = combined
                print(f"[IMAP] Found {len(messages)} messages to process (UNSEEN: {len(unseen_messages)}, Today: {len(today_messages)}).")
            else:
                # Final fallback: last 10 recent messages if inbox is empty today
                all_msgs = client.search('ALL')
                if len(all_msgs) > 10:
                    all_msgs = all_msgs[-10:]
                messages = all_msgs
                print(f"[IMAP] No new messages today. Ingesting latest {len(messages)} recent messages from Inbox...")



            for msg_id in messages:
                response = client.fetch(msg_id, ['ENVELOPE'])
                data = response.get(msg_id)
                if not data or b'ENVELOPE' not in data:
                    continue
                envelope = data[b'ENVELOPE']

                # Parse sender
                sender_addr = "unknown@example.com"
                if envelope.from_:
                    f = envelope.from_[0]
                    mailbox = f.mailbox.decode('utf-8', errors='ignore') if f.mailbox else ""
                    host = f.host.decode('utf-8', errors='ignore') if f.host else ""
                    sender_addr = f"{mailbox}@{host}"

                # Parse subject
                subject = "No Subject"
                if envelope.subject:
                    raw_subj = envelope.subject.decode('utf-8', errors='ignore')
                    try:
                        from email.header import decode_header
                        decoded_parts = decode_header(raw_subj)
                        decoded_str = ""
                        for part, charset in decoded_parts:
                            if isinstance(part, bytes):
                                decoded_str += part.decode(charset or 'utf-8', errors='ignore')
                            else:
                                decoded_str += str(part)
                        subject = decoded_str.strip() or raw_subj
                    except Exception:
                        subject = raw_subj

                # FAST CHECK: Check if this email already exists in our database for this specific user
                existing_id = get_existing_email_id_fast(sender_addr, subject, user_id)
                if existing_id:
                    if msg_id in unseen_messages:
                        client.add_flags(msg_id, '\\Seen')
                    continue

                # ONLY fetch RFC822 for NEW emails that are not yet processed!
                body_resp = client.fetch(msg_id, ['RFC822'])
                raw_bytes = body_resp.get(msg_id, {}).get(b'RFC822', b'')
                body = ""
                try:
                    msg = email_module.message_from_bytes(raw_bytes)
                    if msg.is_multipart():
                        for part in msg.walk():
                            content_type = part.get_content_type()
                            content_disposition = str(part.get("Content-Disposition"))
                            if content_type == "text/plain" and "attachment" not in content_disposition:
                                body = part.get_payload(decode=True).decode('utf-8', errors='ignore')
                                break
                            elif content_type == "text/html" and not body and "attachment" not in content_disposition:
                                html_content = part.get_payload(decode=True).decode('utf-8', errors='ignore')
                                body = re.sub('<[^<]+?>', '', html_content)
                    else:
                        body = msg.get_payload(decode=True).decode('utf-8', errors='ignore')
                except Exception as ex:
                    print(f"[Parser Warning] Error parsing MIME body: {ex}")
                    body = data.get(b'RFC822.TEXT', b'').decode('utf-8', errors='ignore')

                # Fallback if body is empty
                if not body.strip():
                    body = "Empty body."

                insights = None
                if llm_api_key:
                    insights = generate_with_gemini(llm_api_key, sender_addr, subject, body)
                
                if not insights:
                    print(f"[Fallback] Using local intelligent email classifier for '{subject}'")
                    insights = heuristic_fallback(sender_addr, subject, body)

                processed = {
                    "sender": sender_addr,
                    "subject": subject,
                    "body": body[:4000],
                    "priority": insights.get("priority", "Medium"),
                    "category": insights.get("category", "important"),
                    "intent": insights.get("intent", ""),
                    "summary": insights.get("summary", ""),
                    "action_items": insights.get("action_items", []),
                    "draft_reply": insights.get("draft_reply", ""),
                    "requires_approval": 1 if insights.get("requires_approval") else 0
                }
                eid = insert_email_to_db(processed, user_id, imap_account_id)
                if eid:
                    log_to_excel(eid, sender_addr, subject, processed["priority"], processed["category"], processed["intent"], "Pending")
                    if msg_id in unseen_messages:
                        client.add_flags(msg_id, '\\Seen')
                        print(f"[IMAP] Marked email {msg_id} as Read.")

                time.sleep(0.5)

            print(f"[OK] Inbox processing completed for {imap_user}.")

    except Exception as e:
        print(f"[Error for user {imap_user}] {e}")

# ── Database Settings Loader ─────────────────────────────────────────────────

def load_db_settings():
    settings = {}
    if not os.path.exists(DB_FILE):
        return settings
    try:
        conn = sqlite3.connect(DB_FILE)
        c = conn.cursor()
        c.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='settings'")
        if c.fetchone():
            c.execute("SELECT key, value FROM settings")
            for row in c.fetchall():
                if row[1]:
                    settings[row[0]] = row[1]
        conn.close()
    except Exception as e:
        print(f"[DB Warning] {e}")
    return settings

# ── Main Entry Point ─────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="AI Email Automation Background Agent")
    parser.add_argument('--simulate', action='store_true', help="Run simulation with mock emails once and exit")
    parser.add_argument('--loop', action='store_true', help="Run in a continuous 30-second loop")
    parser.add_argument('--user-id', type=int, default=None, help="Target specific user ID")
    args = parser.parse_args()

    load_env()
    target_user_id = args.user_id or 1

    if args.simulate:
        run_simulation(user_id=target_user_id)
        return

    # Helper function to perform a single sync cycle
    def do_sync_cycle():
        db_settings = load_db_settings()

        db_key_enc = db_settings.get("gemini_api_key")
        db_key = decrypt_password(db_key_enc) if db_key_enc else None
        env_key = os.environ.get("GEMINI_API_KEY")
        llm_api_key = db_key if (db_key and not db_key.startswith("test-key") and not db_key.startswith("•••")) else env_key
        env_model = db_settings.get("gemini_model") or os.environ.get("GEMINI_MODEL", "gemini-flash-lite-latest")

        if env_model:
            os.environ["GEMINI_MODEL"] = env_model

        users_to_process = []
        if os.path.exists(DB_FILE):
            try:
                conn = sqlite3.connect(DB_FILE)
                c = conn.cursor()
                c.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='imap_accounts'")
                if c.fetchone():
                    # Check if columns exist
                    c.execute("PRAGMA table_info(imap_accounts)")
                    cols = [col[1] for col in c.fetchall()]
                    has_oauth = 'auth_type' in cols and 'oauth_refresh_token' in cols
                    
                    query = "SELECT id, user_id, email, imap_server, imap_user, imap_password"
                    if has_oauth:
                        query += ", auth_type, oauth_refresh_token"
                    query += " FROM imap_accounts"
                    if args.user_id:
                        query += f" WHERE user_id = {int(args.user_id)}"
                    
                    c.execute(query)
                    for row in c.fetchall():
                        if has_oauth:
                            acc_id, u_id, u_email, u_server, u_user, u_pass, auth_type, oauth_token = row
                            auth_type = auth_type or 'password'
                        else:
                            acc_id, u_id, u_email, u_server, u_user, u_pass = row
                            auth_type = 'password'
                            oauth_token = None

                        if auth_type == 'oauth':
                            oauth_decrypted = decrypt_password(oauth_token) if oauth_token else ""
                            if u_server and u_user and oauth_decrypted:
                                users_to_process.append({
                                    "id": u_id,
                                    "account_id": acc_id,
                                    "email": u_email,
                                    "imap_server": u_server or "imap.gmail.com",
                                    "imap_user": u_user,
                                    "imap_password": "",
                                    "auth_type": "oauth",
                                    "oauth_refresh_token": oauth_decrypted
                                })
                        else:
                            u_pass_decrypted = decrypt_password(u_pass) if u_pass else ""
                            if u_server and u_user and u_pass_decrypted:
                                users_to_process.append({
                                    "id": u_id,
                                    "account_id": acc_id,
                                    "email": u_email,
                                    "imap_server": u_server,
                                    "imap_user": u_user,
                                    "imap_password": u_pass_decrypted,
                                    "auth_type": "password",
                                    "oauth_refresh_token": None
                                })
                conn.close()
            except Exception as e:
                print(f"[DB Error fetching IMAP accounts] {e}")

        # Fallback to global setting if no users in DB have IMAP configured
        if not users_to_process and not args.user_id:
            g_server = db_settings.get("imap_server") or os.environ.get("IMAP_SERVER")
            g_user = db_settings.get("imap_user") or os.environ.get("IMAP_USER")
            g_pass_enc = db_settings.get("imap_password") or os.environ.get("IMAP_PASSWORD")
            g_pass = decrypt_password(g_pass_enc) if g_pass_enc else ""
            if g_server and g_user and g_pass:
                users_to_process.append({
                    "id": 1,
                    "account_id": None,
                    "email": g_user,
                    "imap_server": g_server,
                    "imap_user": g_user,
                    "imap_password": g_pass,
                    "auth_type": "password",
                    "oauth_refresh_token": None
                })

        if not users_to_process:
            print(f"[Info] No active IMAP credentials found for user filter ({args.user_id or 'all'}). Running simulation ingestion...")
            run_simulation(user_id=target_user_id)
        else:
            if not llm_api_key or llm_api_key == "your_gemini_api_key_here":
                print("[Warn] GEMINI_API_KEY is not set. Cannot process real emails without Gemini LLM Key.")
            else:
                # Sync for each user
                for user in users_to_process:
                    run_real_agent(
                        user["imap_server"],
                        user["imap_user"],
                        user["imap_password"],
                        llm_api_key,
                        user["id"],
                        user["account_id"],
                        user.get("auth_type", "password"),
                        user.get("oauth_refresh_token")
                    )

    if args.loop:
        print("AI Email Automation Background Agent started. Sync loop running every 30s...")
        while True:
            try:
                do_sync_cycle()
            except Exception as e:
                print(f"[Agent Loop Error] {e}")
            time.sleep(30)
    else:
        print("AI Email Automation Agent: running sync cycle...")
        do_sync_cycle()

if __name__ == "__main__":
    main()

